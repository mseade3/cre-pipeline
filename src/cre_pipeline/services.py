from __future__ import annotations

import csv
import hashlib
import json
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, TypedDict, cast

import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session, selectinload

from cre_pipeline import ATTRIBUTION
from cre_pipeline.adapters import EntityReviewCsvAdapter, ParcelCsvAdapter
from cre_pipeline.models import (
    Activity,
    Contact,
    DuplicateCandidate,
    Entity,
    FieldProvenance,
    ManualEntityReview,
    Opportunity,
    Ownership,
    Property,
    ScoreComponent,
    ScoringRun,
    Signal,
    Suppression,
)
from cre_pipeline.schemas import ParcelImport, normalize_text

ENTITY_LOOKUP_URL = "https://egov.maryland.gov/BusinessExpress/EntitySearch"
ENTITY_REVIEW_SOURCE = "Maryland Business Express manual review"
RESIDENT_AGENT_WARNING = "Resident agent may not be the owner or beneficial owner."


@dataclass(frozen=True)
class ImportResult:
    inserted: int
    updated: int


@dataclass(frozen=True)
class ComponentResult:
    component: str
    state: str
    weight: float
    contribution: float
    reason: str
    evidence_reference: str | None = None


@dataclass(frozen=True)
class PriorityResult:
    score: float
    coverage: float
    outreach_eligible: bool
    components: list[ComponentResult]


class ScoreConfig(TypedDict):
    version: str
    weights: dict[str, float]
    thresholds: dict[str, float]
    penalties: dict[str, float]


class ScoreInputs(TypedDict, total=False):
    buyer_fit: float | None
    ownership_tenure_years: float | None
    absentee_ownership: bool | None
    vacancy_underutilization: bool | None
    public_records: bool | None
    entity_administrative_status: bool | None
    completeness: float
    freshness: float
    recent_outreach: bool
    suppressed: bool


PROVENANCE_FIELDS = (
    "parcel_id",
    "account_id",
    "address",
    "municipality",
    "postal_code",
    "latitude",
    "longitude",
    "property_type",
    "land_use_code",
    "zoning",
    "land_area_sqft",
    "building_area_sqft",
    "year_built",
    "assessed_land_value",
    "assessed_improvement_value",
    "last_sale_date",
    "last_sale_price",
)


def _property_values(record: ParcelImport, source_name: str) -> dict[str, Any]:
    return {
        "parcel_id": record.parcel_id,
        "account_id": record.account_id,
        "source_name": source_name,
        "source_record_id": record.source_record_id,
        "address": record.address,
        "address_normalized": normalize_text(record.address),
        "municipality": record.municipality,
        "postal_code": record.postal_code,
        "latitude": record.latitude,
        "longitude": record.longitude,
        "property_type": record.property_type,
        "land_use_code": record.land_use_code,
        "zoning": record.zoning,
        "land_area_sqft": record.land_area_sqft,
        "building_area_sqft": record.building_area_sqft,
        "year_built": record.year_built,
        "assessed_land_value": record.assessed_land_value,
        "assessed_improvement_value": record.assessed_improvement_value,
        "last_sale_date": record.last_sale_date,
        "last_sale_price": record.last_sale_price,
        "retrieved_at": record.retrieved_at,
    }


def _upsert_provenance(
    session: Session,
    property_id: int,
    record: ParcelImport,
    source_name: str,
    source_url: str,
) -> None:
    raw = record.model_dump()
    session.execute(
        delete(FieldProvenance).where(
            FieldProvenance.record_type == "property",
            FieldProvenance.record_id == property_id,
            FieldProvenance.source_name == source_name,
        )
    )
    for field_name in PROVENANCE_FIELDS:
        value = raw[field_name]
        source_value = record.source_raw.get(field_name, value)
        session.add(
            FieldProvenance(
                record_type="property",
                record_id=property_id,
                field_name=field_name,
                source_name=source_name,
                source_url=source_url,
                retrieved_at=record.retrieved_at,
                raw_value=(
                    None
                    if source_value is None
                    else json.dumps(source_value, sort_keys=True)
                    if isinstance(source_value, dict)
                    else str(source_value)
                ),
                normalized_value=(
                    normalize_text(value)
                    if isinstance(value, str) and field_name in {"address"}
                    else None
                    if value is None
                    else str(value)
                ),
                confidence=record.confidence,
                evidence_note=record.evidence_note,
            )
        )


def _replace_signals(
    session: Session, property_id: int, record: ParcelImport, source_url: str
) -> None:
    values: dict[str, object | None] = {
        "ownership_tenure_years": record.ownership_tenure_years,
        "buyer_fit": record.buyer_fit,
        "vacancy_underutilization": record.vacancy_evidence,
        "public_records": record.public_record_signal,
        "entity_administrative_status": record.entity_administrative_issue,
    }
    session.execute(delete(Signal).where(Signal.property_id == property_id))
    for signal_type, value in values.items():
        if value is not None:
            session.add(
                Signal(
                    property_id=property_id,
                    signal_type=signal_type,
                    value=str(value).lower() if isinstance(value, bool) else str(value),
                    evidence=record.evidence_note,
                    confidence=record.confidence,
                    observed_at=record.retrieved_at,
                    source_url=source_url,
                )
            )


def ingest_parcels(session: Session, path: Path, source_name: str, source_url: str) -> ImportResult:
    return ingest_records(
        session,
        ParcelCsvAdapter(path).records(),
        source_name=source_name,
        source_url=source_url,
    )


def ingest_records(
    session: Session,
    records: Iterable[ParcelImport],
    source_name: str,
    source_url: str,
) -> ImportResult:
    inserted = 0
    updated = 0
    for record in records:
        property_record = session.scalar(
            select(Property).where(
                Property.source_name == source_name,
                Property.source_record_id == record.source_record_id,
            )
        )
        values = _property_values(record, source_name)
        if property_record is None:
            property_record = Property(**values)
            session.add(property_record)
            session.flush()
            inserted += 1
        else:
            for field_name, value in values.items():
                setattr(property_record, field_name, value)
            session.execute(delete(Ownership).where(Ownership.property_id == property_record.id))
            updated += 1

        session.add(
            Ownership(
                property_id=property_record.id,
                legal_owner_name=record.owner_name,
                owner_name_normalized=normalize_text(record.owner_name),
                owner_type=record.owner_type,
                mailing_address=record.owner_mailing_address,
                absentee_indicator=record.absentee_indicator,
                entity_match_confidence=None,
                ownership_evidence=record.evidence_note,
                verification_status="source_reported",
            )
        )
        _replace_signals(session, property_record.id, record, source_url)
        _upsert_provenance(session, property_record.id, record, source_name, source_url)
    return ImportResult(inserted=inserted, updated=updated)


def import_entity_reviews(session: Session, path: Path) -> ImportResult:
    inserted = 0
    updated = 0
    for review in EntityReviewCsvAdapter(path).records():
        entity = None
        if review.maryland_department_id:
            entity = session.scalar(
                select(Entity).where(Entity.maryland_department_id == review.maryland_department_id)
            )
        if entity is None:
            entity = session.scalar(
                select(Entity).where(Entity.legal_entity_name == review.entity_name)
            )
        if entity is None:
            entity = Entity(legal_entity_name=review.entity_name)
            session.add(entity)
            inserted += 1
        else:
            updated += 1
        entity.maryland_department_id = review.maryland_department_id
        entity.entity_status = review.entity_status
        entity.good_standing_status = review.good_standing_status
        entity.principal_office = review.principal_office
        entity.resident_agent = review.resident_agent
        entity.resident_agent_warning = RESIDENT_AGENT_WARNING
        entity.verification_date = review.review_date
        entity.source_url = review.official_lookup_url
        session.flush()

        normalized_entity_name = normalize_text(review.entity_name)
        matching_queue_items = [
            item
            for item in session.scalars(
                select(ManualEntityReview).order_by(ManualEntityReview.id)
            ).all()
            if normalize_text(item.entity_name) == normalized_entity_name
        ]
        queue_item = next(
            (item for item in matching_queue_items if item.status == "pending"),
            None,
        )
        if queue_item is None:
            queue_item = next(
                (
                    item
                    for item in matching_queue_items
                    if item.status == "completed"
                    and item.reviewer == review.reviewer
                    and item.review_date == review.review_date
                ),
                None,
            )
        if queue_item is None:
            queue_item = ManualEntityReview(
                entity_name=review.entity_name,
                official_lookup_url=review.official_lookup_url,
                requested_fields=(
                    "entity_status,good_standing_status,principal_office,resident_agent"
                ),
            )
            session.add(queue_item)
        queue_item.reviewer = review.reviewer
        queue_item.review_date = review.review_date
        queue_item.evidence_notes = review.evidence_notes
        queue_item.status = "completed"
        session.execute(
            delete(FieldProvenance).where(
                FieldProvenance.record_type == "entity",
                FieldProvenance.record_id == entity.id,
                FieldProvenance.source_name == ENTITY_REVIEW_SOURCE,
            )
        )
        entity_values = {
            "legal_entity_name": review.entity_name,
            "maryland_department_id": review.maryland_department_id,
            "entity_status": review.entity_status,
            "good_standing_status": review.good_standing_status,
            "principal_office": review.principal_office,
            "resident_agent": review.resident_agent,
            "verification_date": review.review_date,
        }
        retrieved_at = datetime.combine(review.review_date, datetime.min.time(), tzinfo=UTC)
        for field_name, value in entity_values.items():
            session.add(
                FieldProvenance(
                    record_type="entity",
                    record_id=entity.id,
                    field_name=field_name,
                    source_name=ENTITY_REVIEW_SOURCE,
                    source_url=review.official_lookup_url,
                    retrieved_at=retrieved_at,
                    raw_value=None if value is None else str(value),
                    normalized_value=(
                        normalize_text(value)
                        if isinstance(value, str)
                        else None
                        if value is None
                        else str(value)
                    ),
                    confidence=review.verification_confidence,
                    evidence_note=review.evidence_notes,
                )
            )

        matching_ownerships = session.scalars(
            select(Ownership).where(Ownership.owner_name_normalized == normalized_entity_name)
        ).all()
        for ownership in matching_ownerships:
            ownership.entity_id = entity.id
            ownership.entity_match_confidence = 1.0
            ownership.verification_status = "manual_entity_name_match"
    return ImportResult(inserted=inserted, updated=updated)


def find_duplicate_candidates(session: Session) -> int:
    session.execute(delete(DuplicateCandidate).where(DuplicateCandidate.status == "pending"))
    properties = session.scalars(
        select(Property).options(selectinload(Property.ownerships)).order_by(Property.id)
    ).all()
    seen: set[tuple[int, int, str]] = set()
    for index, left in enumerate(properties):
        left_owner = left.ownerships[0].owner_name_normalized if left.ownerships else ""
        for right in properties[index + 1 :]:
            right_owner = right.ownerships[0].owner_name_normalized if right.ownerships else ""
            rule: str | None = None
            if left.parcel_id and left.parcel_id == right.parcel_id:
                rule = "parcel_id_exact"
            elif left.account_id and left.account_id == right.account_id:
                rule = "account_id_exact"
            elif (
                left.address_normalized == right.address_normalized
                and left_owner
                and left_owner == right_owner
            ):
                rule = "address_owner_exact"
            if rule:
                key = (left.id, right.id, rule)
                if key not in seen:
                    session.add(
                        DuplicateCandidate(
                            property_id_a=left.id,
                            property_id_b=right.id,
                            match_rule=rule,
                            details=(
                                "Exact deterministic match; human approval required before merge."
                            ),
                        )
                    )
                    seen.add(key)
    return len(seen)


def load_score_config(path: Path) -> ScoreConfig:
    with path.open(encoding="utf-8") as handle:
        raw = yaml.safe_load(handle)
    if not isinstance(raw, dict):
        raise ValueError("scoring configuration must be a mapping")
    config = cast(ScoreConfig, raw)
    if sum(config["weights"].values()) != 100:
        raise ValueError("positive scoring weights must total 100")
    if config["weights"]["entity_administrative_status"] > 3:
        raise ValueError("entity administrative status weight may not exceed 3")
    return config


def calculate_priority(config: ScoreConfig, inputs: ScoreInputs) -> PriorityResult:
    weights = config["weights"]
    threshold = config["thresholds"]["long_tenure_years"]
    components: list[ComponentResult] = []
    known = 0
    total_evidence_components = 6

    def add_binary(key: str, value: bool | None, positive_reason: str) -> None:
        nonlocal known
        if value is None:
            components.append(
                ComponentResult(key, "unknown", weights[key], 0, "Missing data; no inference made.")
            )
            return
        known += 1
        components.append(
            ComponentResult(
                key,
                "supported" if value else "not_supported",
                weights[key],
                weights[key] if value else 0,
                positive_reason if value else "Available evidence does not support this signal.",
            )
        )

    buyer_fit = inputs.get("buyer_fit")
    if buyer_fit is None:
        components.append(
            ComponentResult(
                "buyer_fit", "unknown", weights["buyer_fit"], 0, "Buyer criteria not documented."
            )
        )
    else:
        known += 1
        components.append(
            ComponentResult(
                "buyer_fit",
                "measured",
                weights["buyer_fit"],
                weights["buyer_fit"] * max(0, min(1, buyer_fit)),
                "Contribution reflects documented buyer-criteria fit.",
            )
        )

    tenure = inputs.get("ownership_tenure_years")
    if tenure is None:
        components.append(
            ComponentResult(
                "ownership_tenure",
                "unknown",
                weights["ownership_tenure"],
                0,
                "Ownership start date is unavailable; no inference made.",
            )
        )
    else:
        known += 1
        fraction = min(max(tenure, 0) / threshold, 1)
        components.append(
            ComponentResult(
                "ownership_tenure",
                "measured",
                weights["ownership_tenure"],
                weights["ownership_tenure"] * fraction,
                f"Public/source-reported tenure is {tenure:g} years.",
            )
        )

    add_binary(
        "absentee_ownership",
        inputs.get("absentee_ownership"),
        "Owner mailing address differs from property address.",
    )
    add_binary(
        "vacancy_underutilization",
        inputs.get("vacancy_underutilization"),
        "Source contains evidence of possible vacancy or underutilization; human review required.",
    )
    add_binary(
        "public_records",
        inputs.get("public_records"),
        (
            "Permitted public record contains a relevant research signal; "
            "it does not prove motivation."
        ),
    )
    add_binary(
        "entity_administrative_status",
        inputs.get("entity_administrative_status"),
        "Administrative entity status merits verification and does not prove distress.",
    )

    completeness = max(0, min(1, inputs.get("completeness", 0)))
    freshness = max(0, min(1, inputs.get("freshness", 0)))
    quality = weights["data_quality"] * ((completeness + freshness) / 2)
    components.append(
        ComponentResult(
            "data_quality",
            "measured",
            weights["data_quality"],
            quality,
            "Quality bonus reflects field completeness and source freshness only.",
        )
    )

    score = sum(item.contribution for item in components)
    if inputs.get("recent_outreach", False):
        penalty = config["penalties"]["recent_outreach"]
        components.append(
            ComponentResult(
                "recent_outreach_penalty",
                "applied",
                -penalty,
                -penalty,
                "Recent outreach lowers workflow priority.",
            )
        )
        score -= penalty

    suppressed = inputs.get("suppressed", False)
    return PriorityResult(
        score=round(max(0, min(100, score)), 2),
        coverage=round(known / total_evidence_components, 3),
        outreach_eligible=not suppressed,
        components=components,
    )


def _signal_map(property_record: Property) -> dict[str, str]:
    return {signal.signal_type: signal.value for signal in property_record.signals}


def _to_bool(value: str | None) -> bool | None:
    if value is None:
        return None
    return value.lower() == "true"


def score_all(session: Session, config_path: Path) -> int:
    config = load_score_config(config_path)
    properties = session.scalars(
        select(Property)
        .options(selectinload(Property.ownerships), selectinload(Property.signals))
        .order_by(Property.id)
    ).all()
    for property_record in properties:
        signals = _signal_map(property_record)
        ownership = property_record.ownerships[0] if property_record.ownerships else None
        populated = sum(
            value is not None
            for value in (
                property_record.parcel_id,
                property_record.address,
                property_record.property_type,
                property_record.land_area_sqft,
                property_record.building_area_sqft,
                property_record.year_built,
                property_record.assessed_land_value,
                property_record.assessed_improvement_value,
                ownership.legal_owner_name if ownership else None,
                ownership.mailing_address if ownership else None,
            )
        )
        retrieved = property_record.retrieved_at
        if retrieved.tzinfo is None:
            retrieved = retrieved.replace(tzinfo=UTC)
        age_days = max(0, (datetime.now(UTC) - retrieved).days)
        freshness = max(0.0, 1 - age_days / config["thresholds"]["fresh_days"])
        result = calculate_priority(
            config,
            ScoreInputs(
                buyer_fit=(float(signals["buyer_fit"]) if "buyer_fit" in signals else None),
                ownership_tenure_years=(
                    float(signals["ownership_tenure_years"])
                    if "ownership_tenure_years" in signals
                    else None
                ),
                absentee_ownership=ownership.absentee_indicator if ownership else None,
                vacancy_underutilization=_to_bool(signals.get("vacancy_underutilization")),
                public_records=_to_bool(signals.get("public_records")),
                entity_administrative_status=_to_bool(signals.get("entity_administrative_status")),
                completeness=populated / 10,
                freshness=freshness,
                recent_outreach=False,
                suppressed=False,
            ),
        )
        session.execute(delete(ScoringRun).where(ScoringRun.property_id == property_record.id))
        run = ScoringRun(
            property_id=property_record.id,
            score=result.score,
            data_coverage=result.coverage,
            outreach_eligible=result.outreach_eligible,
            config_version=config["version"],
        )
        session.add(run)
        session.flush()
        for component in result.components:
            session.add(
                ScoreComponent(
                    scoring_run_id=run.id,
                    component=component.component,
                    state=component.state,
                    weight=component.weight,
                    contribution=component.contribution,
                    reason=component.reason,
                    evidence_reference=(
                        component.evidence_reference
                        or (
                            f"property:{property_record.id}/component:{component.component}"
                            if component.state != "unknown"
                            else None
                        )
                    ),
                )
            )
    return len(properties)


def create_pending_entity_reviews(session: Session) -> int:
    ownership_names = session.scalars(select(Ownership.legal_owner_name).distinct()).all()
    existing = set(session.scalars(select(ManualEntityReview.entity_name)).all())
    count = 0
    for name in ownership_names:
        normalized = name.upper()
        looks_like_entity = any(
            suffix in normalized
            for suffix in (" LLC", " L.L.C", " INC", " CORP", " LP", " LLP", " LTD")
        )
        if looks_like_entity and name not in existing:
            session.add(
                ManualEntityReview(
                    entity_name=name,
                    official_lookup_url=ENTITY_LOOKUP_URL,
                    requested_fields=(
                        "entity_status,good_standing_status,principal_office,resident_agent"
                    ),
                    status="pending",
                )
            )
            count += 1
    return count


def export_entity_review_queue(session: Session, output: Path) -> int:
    queue = session.scalars(
        select(ManualEntityReview)
        .where(ManualEntityReview.status == "pending")
        .order_by(ManualEntityReview.entity_name, ManualEntityReview.id)
    ).all()
    output.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "entity_name",
        "maryland_department_id",
        "entity_status",
        "good_standing_status",
        "principal_office",
        "resident_agent",
        "reviewer",
        "review_date",
        "official_lookup_url",
        "evidence_notes",
        "verification_confidence",
        "requested_fields",
        "resident_agent_warning",
    ]
    with output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for item in queue:
            writer.writerow(
                {
                    "entity_name": item.entity_name,
                    "maryland_department_id": "",
                    "entity_status": "",
                    "good_standing_status": "",
                    "principal_office": "",
                    "resident_agent": "",
                    "reviewer": "",
                    "review_date": "",
                    "official_lookup_url": item.official_lookup_url,
                    "evidence_notes": "",
                    "verification_confidence": "0.9",
                    "requested_fields": item.requested_fields,
                    "resident_agent_warning": RESIDENT_AGENT_WARNING,
                }
            )
    return len(queue)


def report_metrics(session: Session) -> dict[str, int | float]:
    def count(model: type[Any]) -> int:
        return int(session.scalar(select(func.count()).select_from(model)) or 0)

    def activity_outcome_count(code: str) -> int:
        return int(
            session.scalar(
                select(func.count()).select_from(Activity).where(Activity.outcome_code == code)
            )
            or 0
        )

    property_count = count(Property)
    duplicate_count = int(
        session.scalar(
            select(func.count())
            .select_from(DuplicateCandidate)
            .where(DuplicateCandidate.status == "pending")
        )
        or 0
    )
    entity_review_count = int(
        session.scalar(
            select(func.count())
            .select_from(ManualEntityReview)
            .where(ManualEntityReview.status == "pending")
        )
        or 0
    )
    entity_review_completed = int(
        session.scalar(
            select(func.count())
            .select_from(ManualEntityReview)
            .where(ManualEntityReview.status == "completed")
        )
        or 0
    )
    return {
        "records_imported": property_count,
        "records_requiring_review": entity_review_count + duplicate_count,
        "entity_reviews_pending": entity_review_count,
        "entity_reviews_completed": entity_review_completed,
        "duplicate_candidates": duplicate_count,
        "duplicate_rate": round(duplicate_count / property_count, 3) if property_count else 0.0,
        "contacts": count(Contact),
        "activities": count(Activity),
        "calls_attempted": int(
            session.scalar(
                select(func.count()).select_from(Activity).where(Activity.activity_type == "call")
            )
            or 0
        ),
        "conversations": activity_outcome_count("conversation"),
        "decision_makers_reached": activity_outcome_count("decision_maker_reached"),
        "follow_ups": int(
            session.scalar(
                select(func.count()).select_from(Activity).where(Activity.follow_up_at.is_not(None))
            )
            or 0
        ),
        "seller_stated_interest": activity_outcome_count("seller_stated_interest"),
        "opportunities_referred": activity_outcome_count("referred_to_partner"),
        "offers": activity_outcome_count("offer"),
        "closed_transactions": activity_outcome_count("closed"),
        "opportunities": count(Opportunity),
        "opt_outs_and_suppressions": count(Suppression),
    }


SHEET_NAMES = (
    "Qualified Leads",
    "Research Queue",
    "Outreach Queue",
    "Follow-Ups",
    "Suppressions",
    "Data Dictionary",
    "Score Explanation",
    "Compliance Checklist",
)


def _latest_scores(session: Session) -> dict[int, ScoringRun]:
    runs = session.scalars(
        select(ScoringRun)
        .options(selectinload(ScoringRun.components))
        .order_by(ScoringRun.property_id, ScoringRun.created_at.desc())
    ).all()
    return {run.property_id: run for run in runs}


def export_workbook(session: Session, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    scores = _latest_scores(session)
    properties = session.scalars(
        select(Property).options(selectinload(Property.ownerships)).order_by(Property.id)
    ).all()
    lead_rows: list[dict[str, object]] = []
    score_rows: list[dict[str, object]] = []
    for prop in properties:
        owner = prop.ownerships[0] if prop.ownerships else None
        run = scores.get(prop.id)
        lead_rows.append(
            {
                "property_id": prop.id,
                "parcel_id": prop.parcel_id,
                "account_id": prop.account_id,
                "address": prop.address,
                "property_type": prop.property_type,
                "owner_name": owner.legal_owner_name if owner else None,
                "verification_status": owner.verification_status if owner else "missing",
                "research_priority": run.score if run else None,
                "data_coverage": run.data_coverage if run else None,
                "outreach_eligible": run.outreach_eligible if run else False,
                "human_approval": "Not approved",
                "source": prop.source_name,
                "retrieved_at": prop.retrieved_at.isoformat(),
                "attribution": ATTRIBUTION,
            }
        )
        if run:
            for component in run.components:
                score_rows.append(
                    {
                        "property_id": prop.id,
                        "component": component.component,
                        "state": component.state,
                        "weight": component.weight,
                        "contribution": component.contribution,
                        "reason": component.reason,
                        "evidence_reference": component.evidence_reference,
                        "config_version": run.config_version,
                    }
                )

    duplicates = session.execute(
        select(
            DuplicateCandidate.id,
            DuplicateCandidate.property_id_a,
            DuplicateCandidate.property_id_b,
            DuplicateCandidate.match_rule,
            DuplicateCandidate.status,
            DuplicateCandidate.details,
        )
    ).all()
    entity_reviews = session.execute(
        select(
            ManualEntityReview.id,
            ManualEntityReview.entity_name,
            ManualEntityReview.official_lookup_url,
            ManualEntityReview.requested_fields,
            ManualEntityReview.reviewer,
            ManualEntityReview.review_date,
            ManualEntityReview.status,
            ManualEntityReview.evidence_notes,
        )
    ).all()
    research_rows = [dict(row._mapping) for row in duplicates] + [
        {"review_type": "entity", **dict(row._mapping)} for row in entity_reviews
    ]
    suppressions = session.execute(
        select(
            Suppression.channel,
            Suppression.normalized_value_hash,
            Suppression.reason,
            Suppression.requested_at,
            Suppression.source,
            Suppression.active,
        )
    ).all()
    follow_ups = session.execute(
        select(
            Activity.property_id,
            Activity.contact_id,
            Activity.activity_type,
            Activity.outcome_code,
            Activity.outcome,
            Activity.follow_up_at,
            Activity.user_name,
            Contact.outreach_approved,
            Contact.do_not_contact,
            Contact.opt_out_state,
        )
        .outerjoin(Contact, Contact.id == Activity.contact_id)
        .where(Activity.follow_up_at.is_not(None))
        .order_by(Activity.follow_up_at)
    ).all()

    data_dictionary = [
        {
            "field": "research_priority",
            "definition": "Explainable 0–100 research workflow priority; not a distress score.",
        },
        {
            "field": "outreach_eligible",
            "definition": "False when controls block outreach; never constitutes approval.",
        },
        {
            "field": "human_approval",
            "definition": "Must be explicitly approved outside this export before communication.",
        },
        {
            "field": "resident_agent",
            "definition": "Service-of-process contact; may not be the owner or beneficial owner.",
        },
        {"field": "attribution", "definition": ATTRIBUTION},
    ]
    checklist = [
        {"control": line.strip()[4:], "status": "Attorney/human review required"}
        for line in Path("docs/COMPLIANCE_CHECKLIST.md").read_text(encoding="utf-8").splitlines()
        if line.startswith("- [ ]")
    ]

    frames = {
        "Qualified Leads": pd.DataFrame(lead_rows),
        "Research Queue": pd.DataFrame(research_rows),
        "Outreach Queue": pd.DataFrame(
            [
                {
                    **row,
                    "queue_status": "Blocked pending human approval",
                }
                for row in lead_rows
                if row["outreach_eligible"]
            ]
        ),
        "Follow-Ups": pd.DataFrame([dict(row._mapping) for row in follow_ups]),
        "Suppressions": pd.DataFrame([dict(row._mapping) for row in suppressions]),
        "Data Dictionary": pd.DataFrame(data_dictionary),
        "Score Explanation": pd.DataFrame(score_rows),
        "Compliance Checklist": pd.DataFrame(checklist),
    }
    for name, frame in frames.items():
        if frame.empty:
            frames[name] = pd.DataFrame({"status": ["No records"]})

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name in SHEET_NAMES:
            frames[name].to_excel(writer, sheet_name=name, index=False)

    workbook = load_workbook(output)
    workbook.properties.title = "CRE Research Pipeline Export"
    workbook.properties.subject = (
        f"Research workflow only; not an appraisal. Source attribution: {ATTRIBUTION}"
    )
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True)
        for column in worksheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
            column_index = cast(int, column[0].column)
            worksheet.column_dimensions[get_column_letter(column_index)].width = width
    leads_sheet = workbook["Qualified Leads"]
    if leads_sheet.max_row > 1:
        headers = {cell.value: cell.column for cell in leads_sheet[1]}
        score_column = get_column_letter(cast(int, headers["research_priority"]))
        leads_sheet.conditional_formatting.add(
            f"{score_column}2:{score_column}{leads_sheet.max_row}",
            ColorScaleRule(  # type: ignore[no-untyped-call]
                start_type="min",
                start_color="E7E6E6",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFF2CC",
                end_type="max",
                end_color="BDD7EE",
            ),
        )
        approval_column = get_column_letter(cast(int, headers["human_approval"]))
        validation = DataValidation(
            type="list", formula1='"Not approved,Approved by human,Do not contact"'
        )
        leads_sheet.add_data_validation(validation)
        validation.add(f"{approval_column}2:{approval_column}{leads_sheet.max_row}")
    workbook.save(output)


def value_hash(value: str) -> str:
    return hashlib.sha256(normalize_text(value).encode()).hexdigest()
