from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import Engine, func, select
from sqlalchemy.orm import Session

from cre_pipeline import ATTRIBUTION
from cre_pipeline.models import DuplicateCandidate, FieldProvenance, Property, ScoringRun
from cre_pipeline.services import (
    SHEET_NAMES,
    export_workbook,
    find_duplicate_candidates,
    ingest_parcels,
    score_all,
)

FIXTURE = Path("tests/fixtures/parcels.csv")


def test_import_is_idempotent_and_preserves_provenance(session: Session) -> None:
    first = ingest_parcels(session, FIXTURE, "fixture", "file://fixture.csv")
    session.commit()
    second = ingest_parcels(session, FIXTURE, "fixture", "file://fixture.csv")
    session.commit()

    assert first.inserted == 3
    assert second.updated == 3
    assert session.scalar(select(func.count()).select_from(Property)) == 3
    assert session.scalar(select(func.count()).select_from(FieldProvenance)) > 30


def test_deduplication_only_queues_exact_candidates(session: Session) -> None:
    ingest_parcels(session, FIXTURE, "fixture", "file://fixture.csv")
    session.commit()
    assert find_duplicate_candidates(session) == 1
    session.commit()
    candidate = session.scalar(select(DuplicateCandidate))
    assert candidate is not None
    assert candidate.match_rule == "address_owner_exact"
    assert candidate.status == "pending"


def test_full_fixture_pipeline_and_workbook(engine: Engine, tmp_path: Path) -> None:
    with Session(engine) as session, session.begin():
        ingest_parcels(session, FIXTURE, "fixture", "file://fixture.csv")
        find_duplicate_candidates(session)
        assert score_all(session, Path("config/scoring.yaml")) == 3

    output = tmp_path / "leads.xlsx"
    with Session(engine) as session:
        export_workbook(session, output)
        assert session.scalar(select(func.count()).select_from(ScoringRun)) == 3

    workbook = load_workbook(output)
    assert tuple(workbook.sheetnames) == SHEET_NAMES
    for worksheet in workbook.worksheets:
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref is not None
    assert workbook["Qualified Leads"]["N2"].value == ATTRIBUTION
    assert len(workbook["Qualified Leads"].data_validations.dataValidation) == 1
